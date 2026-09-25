"""Import existing session workbooks and add professional beam plots."""

from __future__ import annotations

import math
import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

try:
    from .session_excel import (
        BeamAnalysis,
        BeamInterpolationGrid,
        SessionExcelReport,
        SessionMeasurementRecord,
        interpolate_beam_grid,
    )
except ImportError:  # Allows direct execution from the gui directory.
    from session_excel import (  # type: ignore
        BeamAnalysis,
        BeamInterpolationGrid,
        SessionExcelReport,
        SessionMeasurementRecord,
        interpolate_beam_grid,
    )


DEFAULT_X_PULSES_PER_MM = 160.0
DEFAULT_Z_PULSES_PER_MM = 91.4285714286
DEFAULT_TARGET_DIAMETER_MM = 100.0

_IDLE_POSITION_PATTERN = re.compile(r"\bX=(-?\d+)\s+Z=(-?\d+)\b", re.IGNORECASE)


@dataclass(frozen=True)
class BeamWorkbookDataset:
    source_path: Path
    report: SessionExcelReport
    analysis: BeamAnalysis
    weather_grid: BeamInterpolationGrid
    sheet_name: str
    measurement_rows: int
    valid_rows: int
    pulse_position_rows: int
    x_header: str
    z_header: str
    power_header: str
    x_pulses_per_mm: float
    z_pulses_per_mm: float

    @property
    def position_source_summary(self) -> str:
        if self.pulse_position_rows == self.valid_rows and self.valid_rows:
            return "EVENT IDLE pulse positions"
        if self.pulse_position_rows:
            return (
                f"EVENT IDLE pulses for {self.pulse_position_rows:,} rows; "
                "worksheet X/Z values for the remainder"
            )
        return "worksheet X/Z columns"


@dataclass(frozen=True)
class BeamPlotExportResult:
    source_path: Path
    output_path: Path
    valid_rows: int
    inside_rows: int
    peak_power_dbm: float | None
    centroid_offset_mm: float | None


def load_beam_workbook(
    source_path: str | Path,
    *,
    target_diameter_mm: float | None = None,
    x_pulses_per_mm: float | None = None,
    z_pulses_per_mm: float | None = None,
    use_event_positions: bool = True,
    center_x_mm: float | None = None,
    center_z_mm: float | None = None,
) -> BeamWorkbookDataset:
    """Read a legacy or current session workbook without modifying it."""
    source = Path(source_path).resolve()
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Select an Excel .xlsx session workbook")
    if not source.is_file():
        raise FileNotFoundError(f"Excel workbook not found: {source}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        metadata = _read_summary_metadata(workbook)
        worksheet, header_row, headers = _find_measurement_sheet(workbook)
        target_diameter = _positive_value(
            target_diameter_mm,
            metadata.get("Target diameter (mm)"),
            DEFAULT_TARGET_DIAMETER_MM,
            label="Target diameter",
        )
        x_scale = _positive_value(
            x_pulses_per_mm,
            metadata.get("X pulses/mm"),
            DEFAULT_X_PULSES_PER_MM,
            label="X pulses/mm",
        )
        z_scale = _positive_value(
            z_pulses_per_mm,
            metadata.get("Z pulses/mm"),
            DEFAULT_Z_PULSES_PER_MM,
            label="Z pulses/mm",
        )
        if target_diameter > 2000.0:
            raise ValueError("Target diameter is unreasonably large; enter millimetres")

        x_index, x_header = _required_header(headers, _X_HEADERS, "X position")
        z_index, z_header = _required_header(headers, _Z_HEADERS, "Z position")
        power_index, power_header = _required_header(headers, _POWER_HEADERS, "Power (dBm)")
        result_index = _optional_header(headers, _RESULT_HEADERS)
        notes_index = _optional_header(headers, _NOTES_HEADERS)
        action_index = _optional_header(headers, _ACTION_HEADERS)

        started_at = _as_datetime(metadata.get("Started")) or datetime.now()
        report_metadata = dict(metadata)
        report_metadata["Target diameter (mm)"] = target_diameter
        report_metadata["X pulses/mm"] = x_scale
        report_metadata["Z pulses/mm"] = z_scale
        if center_x_mm is not None:
            report_metadata["Target center X (mm)"] = _finite(center_x_mm, "Target center X")
        if center_z_mm is not None:
            report_metadata["Target center Z (mm)"] = _finite(center_z_mm, "Target center Z")

        session_id = str(metadata.get("Session ID") or source.stem)
        report = SessionExcelReport(
            source.with_name(f"{source.stem}_BeamPlots.xlsx"),
            session_id=session_id,
            started_at=started_at,
            metadata=report_metadata,
        )

        measurement_rows = 0
        pulse_position_rows = 0
        for worksheet_row, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not any(value is not None for value in row):
                continue
            measurement_rows += 1
            result = str(_row_value(row, result_index) or "OK").strip()
            if result and result.upper() != "OK":
                continue
            try:
                x_mm = _finite(_row_value(row, x_index), f"X position in row {worksheet_row}")
                z_mm = _finite(_row_value(row, z_index), f"Z position in row {worksheet_row}")
                power_dbm = _finite(
                    _row_value(row, power_index),
                    f"Power in row {worksheet_row}",
                )
            except ValueError:
                continue

            notes = str(_row_value(row, notes_index) or "")
            if use_event_positions and notes:
                match = _IDLE_POSITION_PATTERN.search(notes)
                if match:
                    x_mm = int(match.group(1)) / x_scale
                    z_mm = int(match.group(2)) / z_scale
                    pulse_position_rows += 1

            action_value = _row_value(row, action_index)
            try:
                action_number = int(action_value)
            except (TypeError, ValueError):
                action_number = worksheet_row - header_row
            report.add_measurement(
                _analysis_record(
                    action_number=action_number,
                    x_mm=x_mm,
                    z_mm=z_mm,
                    power_dbm=power_dbm,
                    result="OK",
                    notes=notes,
                    timestamp=started_at,
                )
            )
    finally:
        workbook.close()

    if not report.measurements:
        raise ValueError(
            "No valid rows were found. The workbook needs numeric X position, "
            "Z position and Power (dBm) columns."
        )
    analysis = report.analyze_beam()
    return BeamWorkbookDataset(
        source_path=source,
        report=report,
        analysis=analysis,
        weather_grid=interpolate_beam_grid(analysis),
        sheet_name=worksheet.title,
        measurement_rows=measurement_rows,
        valid_rows=len(report.measurements),
        pulse_position_rows=pulse_position_rows,
        x_header=x_header,
        z_header=z_header,
        power_header=power_header,
        x_pulses_per_mm=x_scale,
        z_pulses_per_mm=z_scale,
    )


def export_plotted_workbook(
    dataset: BeamWorkbookDataset,
    output_path: str | Path | None = None,
) -> BeamPlotExportResult:
    """Save a plotted copy while keeping the selected source workbook unchanged."""
    output = Path(output_path) if output_path is not None else suggested_output_path(dataset.source_path)
    output = output.resolve()
    if output.suffix.lower() != ".xlsx":
        output = output.with_suffix(".xlsx")
    if output == dataset.source_path:
        raise ValueError("The plotted workbook must not overwrite the selected source workbook")
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(dataset.source_path, data_only=False)
    temporary = output.with_name(f"{output.stem}.tmp{output.suffix}")
    try:
        analysis = dataset.report.add_beam_analysis_sheets(workbook)
        workbook.properties.title = "Laser Beam Spatial Analysis"
        workbook.properties.subject = "Imported X-Z position and HP E5574A power analysis"
        workbook.properties.creator = "Laser X-Z Motion Controller"
        workbook.properties.description = (
            "Original session workbook preserved with regenerated measured, weather-style, "
            "profile, heatmap and beam-data analysis sheets."
        )
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass

    return BeamPlotExportResult(
        source_path=dataset.source_path,
        output_path=output,
        valid_rows=dataset.valid_rows,
        inside_rows=len(analysis.inside_points),
        peak_power_dbm=analysis.peak_power_dbm,
        centroid_offset_mm=analysis.centroid_offset_mm,
    )


def suggested_output_path(source_path: str | Path) -> Path:
    source = Path(source_path).resolve()
    candidate = source.with_name(f"{source.stem}_BeamPlots.xlsx")
    if not candidate.exists():
        return candidate
    return source.with_name(
        f"{source.stem}_BeamPlots_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    )


def _find_measurement_sheet(workbook: Any) -> tuple[Any, int, dict[str, tuple[int, str]]]:
    preferred = ["Measurements", "Beam Data"]
    worksheets = [workbook[name] for name in preferred if name in workbook.sheetnames]
    worksheets.extend(sheet for sheet in workbook.worksheets if sheet not in worksheets)
    for worksheet in worksheets:
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(15, worksheet.max_row), values_only=True),
            start=1,
        ):
            headers = {
                _normalize_header(value): (index, str(value).strip())
                for index, value in enumerate(row)
                if value is not None and str(value).strip()
            }
            if (
                _optional_header(headers, _X_HEADERS) is not None
                and _optional_header(headers, _Z_HEADERS) is not None
                and _optional_header(headers, _POWER_HEADERS) is not None
            ):
                return worksheet, row_number, headers
    raise ValueError(
        "No measurement table was found. Expected X position, Z position and Power (dBm) headers."
    )


def _read_summary_metadata(workbook: Any) -> dict[str, Any]:
    if "Summary" not in workbook.sheetnames:
        return {}
    sheet = workbook["Summary"]
    metadata: dict[str, Any] = {}
    for row in range(1, min(sheet.max_row, 100) + 1):
        for label_column, value_column in ((1, 2), (4, 5)):
            label = sheet.cell(row, label_column).value
            value = sheet.cell(row, value_column).value
            if isinstance(label, str) and label.strip() and value is not None:
                metadata[label.strip()] = value
    return metadata


def _analysis_record(
    *,
    action_number: int,
    x_mm: float,
    z_mm: float,
    power_dbm: float,
    result: str,
    notes: str,
    timestamp: datetime,
) -> SessionMeasurementRecord:
    return SessionMeasurementRecord(
        action_number=action_number,
        measurement_timestamp=timestamp,
        elapsed_s=0.0,
        loop_number=0,
        phase="Imported workbook",
        phase_step=0,
        phase_steps=0,
        command="IMPORTED",
        axis="",
        delta_mm=0.0,
        expected_x_mm=x_mm,
        expected_z_mm=z_mm,
        move_started_at=timestamp,
        motion_completed_at=timestamp,
        measurement_started_at=timestamp,
        measurement_completed_at=timestamp,
        dwell_target_s=0.0,
        actual_post_move_wait_s=0.0,
        power_dbm=power_dbm,
        raw_response=str(power_dbm),
        result=result,
        notes=notes,
    )


def _required_header(
    headers: Mapping[str, tuple[int, str]],
    candidates: Iterable[str],
    label: str,
) -> tuple[int, str]:
    index = _optional_header(headers, candidates)
    if index is None:
        raise ValueError(f"Missing required {label} column")
    normalized = next(candidate for candidate in candidates if candidate in headers)
    return index, headers[normalized][1]


def _optional_header(
    headers: Mapping[str, tuple[int, str]],
    candidates: Iterable[str],
) -> int | None:
    for candidate in candidates:
        if candidate in headers:
            return headers[candidate][0]
    return None


def _normalize_header(value: Any) -> str:
    return " ".join(str(value).strip().lower().replace("²", "^2").split())


def _row_value(row: tuple[Any, ...], index: int | None) -> Any:
    return row[index] if index is not None and index < len(row) else None


def _positive_value(primary: Any, secondary: Any, fallback: float, *, label: str) -> float:
    selected = primary if primary is not None else secondary if secondary not in (None, "") else fallback
    number = _finite(selected, label)
    if number <= 0:
        raise ValueError(f"{label} must be greater than zero")
    return number


def _finite(value: Any, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric") from exc
    if not math.isfinite(number):
        raise ValueError(f"{label} must be finite")
    return number


def _as_datetime(value: Any) -> datetime | None:
    return value if isinstance(value, datetime) else None


_X_HEADERS = (
    "controller x position (mm)",
    "expected x (mm)",
    "x position (mm)",
    "x (mm)",
)
_Z_HEADERS = (
    "controller z position (mm)",
    "expected z (mm)",
    "z position (mm)",
    "z (mm)",
)
_POWER_HEADERS = (
    "power (dbm)",
    "power dbm",
    "dbm",
)
_RESULT_HEADERS = ("result", "status")
_NOTES_HEADERS = ("notes", "event idle", "controller event")
_ACTION_HEADERS = ("action #", "action", "point #")
