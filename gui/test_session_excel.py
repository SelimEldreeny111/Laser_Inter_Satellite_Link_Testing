"""Tests for the generated automated-session Excel report."""

from __future__ import annotations

import unittest
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .session_excel import (
    SessionEventRecord,
    SessionExcelReport,
    SessionMeasurementRecord,
)


class SessionExcelReportTests(unittest.TestCase):
    def test_report_contains_summary_measurement_and_event_tables(self) -> None:
        output = Path.cwd() / "tmp" / "test-session-report.xlsx"
        output.unlink(missing_ok=True)
        self.addCleanup(lambda: output.unlink(missing_ok=True))
        started = datetime(2026, 8, 14, 18, 30, 0)
        report = SessionExcelReport(
            output,
            session_id="20260814-183000-000",
            started_at=started,
            metadata={
                "Recipe": "Stepped Z serpentine scan",
                "Controller firmware": "2.4",
                "VISA resource": "GPIB0::24::INSTR",
                "Instrument": "HEWLETT-PACKARD,E5574A,0,0",
                "Scan axis": "Z",
                "Cross axis": "X",
                "Loops": 1,
                "Scan span (mm)": 2.0,
                "Scan increment (mm)": 1.0,
                "Cross shift (mm)": 1.0,
                "Target diameter (mm)": 100.0,
                "Post-move dwell (s)": 1.0,
                "Motion speed (mm/s)": 5.0,
                "Acceleration (mm/s^2)": 20.0,
                "Expected actions": 6,
            },
        )
        report.add_event(SessionEventRecord(started, 0.0, "SESSION_START", "Started"))
        report.add_measurement(
            SessionMeasurementRecord(
                action_number=1,
                measurement_timestamp=started + timedelta(seconds=1.5),
                elapsed_s=1.5,
                loop_number=1,
                phase="Z+ pass",
                phase_step=1,
                phase_steps=2,
                command="MOVE_MM Z 1",
                axis="Z",
                delta_mm=1.0,
                expected_x_mm=0.0,
                expected_z_mm=1.0,
                move_started_at=started,
                motion_completed_at=started + timedelta(seconds=0.45),
                measurement_started_at=started + timedelta(seconds=0.95),
                measurement_completed_at=started + timedelta(seconds=1.0),
                dwell_target_s=1.0,
                actual_post_move_wait_s=0.55,
                power_dbm=-12.25,
                raw_response="-12.25",
                result="OK",
                notes="EVENT IDLE X=0 Z=6",
            )
        )

        report.save(
            status="Completed",
            reason="All requested loops completed",
            ended_at=started + timedelta(seconds=10),
            duration_s=10.0,
            completed_actions=1,
        )

        workbook = load_workbook(output, data_only=False)
        try:
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Summary",
                    "Beam Map",
                    "Weather Map",
                    "Beam Profiles",
                    "Beam Heatmap",
                    "Beam Data",
                    "Measurements",
                    "Events",
                    "Chart Data",
                ],
            )
            self.assertEqual(workbook["Summary"]["B6"].value, "Completed")
            self.assertIn("$T$2:$T$2", workbook["Summary"]["B24"].value)
            self.assertEqual(workbook["Summary"]["E21"].value, 100.0)
            self.assertEqual(len(workbook["Beam Map"]._charts), 2)
            self.assertEqual(len(workbook["Beam Profiles"]._charts), 2)
            spatial_chart, radial_chart = workbook["Beam Map"]._charts
            self.assertEqual(spatial_chart.x_axis.axPos, "b")
            self.assertEqual(spatial_chart.y_axis.axPos, "l")
            self.assertIn("$H$2", spatial_chart.series[0].xVal.numRef.f)
            self.assertIn("$I$2", spatial_chart.series[0].yVal.numRef.f)
            self.assertIn("$AG$2", radial_chart.series[0].xVal.numRef.f)
            self.assertIn("$AH$2", radial_chart.series[0].yVal.numRef.f)
            self.assertEqual(workbook["Beam Heatmap"]["B5"].value, -12.25)
            self.assertEqual(workbook["Weather Map"]["V26"].value, 0.0)
            self.assertEqual(workbook["Beam Data"]["C2"].value, 0.0)
            self.assertEqual(workbook["Beam Data"]["D2"].value, 1.0)
            self.assertEqual(workbook["Beam Data"]["E2"].value, -12.25)
            self.assertEqual(workbook["Chart Data"].sheet_state, "hidden")
            self.assertEqual(workbook["Measurements"]["T2"].value, -12.25)
            self.assertEqual(workbook["Measurements"]["V2"].value, "OK")
            self.assertEqual(workbook["Measurements"].freeze_panes, "F2")
            self.assertIn("SessionMeasurements", workbook["Measurements"].tables)
            self.assertEqual(workbook["Events"]["D2"].value, "SESSION_START")
            self.assertIn("SessionEvents", workbook["Events"].tables)
        finally:
            workbook.close()

    def test_report_handles_session_with_no_valid_power_rows(self) -> None:
        output = Path.cwd() / "tmp" / "test-empty-session-report.xlsx"
        output.unlink(missing_ok=True)
        self.addCleanup(lambda: output.unlink(missing_ok=True))
        started = datetime(2026, 8, 15, 12, 0, 0)
        report = SessionExcelReport(
            output,
            session_id="EMPTY",
            started_at=started,
            metadata={"Target diameter (mm)": 100.0},
        )
        report.save(
            status="Stopped",
            reason="Stopped before first reading",
            ended_at=started,
            duration_s=0.0,
            completed_actions=0,
        )

        workbook = load_workbook(output, data_only=False)
        try:
            self.assertEqual(len(workbook["Beam Map"]._charts), 0)
            self.assertEqual(len(workbook["Beam Profiles"]._charts), 0)
            self.assertIn("No valid completed power measurements", workbook["Beam Map"]["A18"].value)
            self.assertEqual(workbook["Beam Data"].max_row, 1)
            self.assertEqual(workbook["Chart Data"].sheet_state, "hidden")
        finally:
            workbook.close()

    def test_complete_scan_maps_keep_readings_outside_target_circle(self) -> None:
        output = Path.cwd() / "tmp" / "test-complete-scan-report.xlsx"
        output.unlink(missing_ok=True)
        self.addCleanup(lambda: output.unlink(missing_ok=True))
        started = datetime(2026, 8, 15, 13, 30, 0)
        report = SessionExcelReport(
            output,
            session_id="FULL-SCAN",
            started_at=started,
            metadata={
                "Target diameter (mm)": 100.0,
                "Target center X (mm)": 0.0,
                "Target center Z (mm)": 0.0,
            },
        )
        for action, x_mm, z_mm, power_dbm in (
            (1, 0.0, 0.0, -10.0),
            (2, 100.0, 0.0, -20.0),
            (3, 0.0, 100.0, -30.0),
        ):
            report.add_measurement(
                SessionMeasurementRecord(
                    action_number=action,
                    measurement_timestamp=started,
                    elapsed_s=float(action),
                    loop_number=1,
                    phase="scan",
                    phase_step=action,
                    phase_steps=3,
                    command="MOVE",
                    axis="X",
                    delta_mm=0.0,
                    expected_x_mm=x_mm,
                    expected_z_mm=z_mm,
                    move_started_at=started,
                    motion_completed_at=started,
                    measurement_started_at=started,
                    measurement_completed_at=started,
                    dwell_target_s=1.0,
                    actual_post_move_wait_s=1.0,
                    power_dbm=power_dbm,
                    raw_response=str(power_dbm),
                    result="OK",
                )
            )
        report.save(
            status="Completed",
            reason="Complete",
            ended_at=started,
            duration_s=3.0,
            completed_actions=3,
        )

        workbook = load_workbook(output, data_only=False)
        try:
            self.assertEqual(workbook["Beam Data"]["H2"].value, "Yes")
            self.assertEqual(workbook["Beam Data"]["H3"].value, "No")
            self.assertEqual(workbook["Beam Data"]["H4"].value, "No")
            self.assertEqual(workbook["Beam Heatmap"]["B5"].value, -30.0)
            self.assertEqual(workbook["Beam Heatmap"]["C6"].value, -20.0)
            self.assertIsNotNone(workbook["Weather Map"]["B6"].value)
            self.assertIsNotNone(workbook["Weather Map"]["AP46"].value)
            plotted_band_x_values = [
                workbook["Chart Data"].cell(row, column).value
                for column in range(8, 20, 2)
                for row in range(2, 5)
            ]
            self.assertEqual(sum(value is not None for value in plotted_band_x_values), 3)
            spatial_chart = workbook["Beam Map"]._charts[0]
            self.assertLessEqual(spatial_chart.x_axis.scaling.min, -50.0)
            self.assertGreaterEqual(spatial_chart.x_axis.scaling.max, 100.0)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
