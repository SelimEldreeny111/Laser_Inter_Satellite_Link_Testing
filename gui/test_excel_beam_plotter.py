"""Tests for plotting beam data from existing Excel session workbooks."""

from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .excel_beam_plotter import export_plotted_workbook, load_beam_workbook
from .laser_stage_gui import LaserStageApp


class ExcelBeamPlotterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.source = Path.cwd() / "tmp" / "legacy-session.xlsx"
        self.output = Path.cwd() / "tmp" / "legacy-session-plotted.xlsx"
        self.source.parent.mkdir(parents=True, exist_ok=True)
        self.source.unlink(missing_ok=True)
        self.output.unlink(missing_ok=True)
        self.addCleanup(lambda: self.source.unlink(missing_ok=True))
        self.addCleanup(lambda: self.output.unlink(missing_ok=True))
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Session ID", "LEGACY-TEST", None, "Scan axis", "Z"])
        measurements = workbook.create_sheet("Measurements")
        measurements.append(
            [
                "Action #",
                "Expected X (mm)",
                "Expected Z (mm)",
                "Power (dBm)",
                "Result",
                "Notes",
            ]
        )
        measurements.append([1, 0.0, 1.0, -12.0, "OK", "EVENT IDLE X=0 Z=6"])
        measurements.append([2, 10.0, 1.0, -10.0, "OK", "EVENT IDLE X=100 Z=6"])
        measurements.append([3, 10.0, 11.0, -9.0, "OK", "EVENT IDLE X=100 Z=63"])
        measurements.append([4, 0.0, 11.0, -11.0, "OK", "EVENT IDLE X=0 Z=63"])
        measurements.append([5, 5.0, 5.0, None, "FAILED", "timeout"])
        workbook.create_sheet("Events")
        workbook.save(self.source)
        workbook.close()

    def test_legacy_workbook_uses_idle_pulses_and_exports_charts(self) -> None:
        dataset = load_beam_workbook(
            self.source,
            target_diameter_mm=100.0,
            x_pulses_per_mm=10.0,
            z_pulses_per_mm=5.7143,
        )
        self.assertEqual(dataset.valid_rows, 4)
        self.assertEqual(dataset.pulse_position_rows, 4)
        self.assertAlmostEqual(
            dataset.report.measurements[0].expected_z_mm,
            6.0 / 5.7143,
            places=6,
        )

        result = export_plotted_workbook(dataset, self.output)
        self.assertEqual(result.output_path, self.output.resolve())
        source_workbook = load_workbook(self.source, read_only=True)
        plotted_workbook = load_workbook(self.output, data_only=False)
        try:
            self.assertEqual(source_workbook.sheetnames, ["Summary", "Measurements", "Events"])
            self.assertEqual(
                plotted_workbook.sheetnames,
                [
                    "Summary",
                    "Beam Map",
                    "Weather Map",
                    "Beam Profiles",
                    "Beam Heatmap",
                    "Beam Data",
                    "Measurements",
                    "Events",
                    "Chart Data",
                ],
            )
            self.assertEqual(len(plotted_workbook["Beam Map"]._charts), 2)
            self.assertEqual(len(plotted_workbook["Beam Profiles"]._charts), 2)
            self.assertEqual(plotted_workbook["Chart Data"].sheet_state, "hidden")
            self.assertEqual(dataset.weather_grid.size, 41)
            self.assertEqual(plotted_workbook["Beam Data"]["C2"].value, 0.0)
            self.assertAlmostEqual(
                plotted_workbook["Beam Data"]["D2"].value,
                6.0 / 5.7143,
                places=6,
            )
        finally:
            source_workbook.close()
            plotted_workbook.close()

    def test_source_must_contain_position_and_power_headers(self) -> None:
        workbook = Workbook()
        workbook.active["A1"] = "Unrelated data"
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "No measurement table"):
            load_beam_workbook(self.source)

    def test_gui_contains_offline_plotter_tab_and_draws_preview(self) -> None:
        dataset = load_beam_workbook(self.source)
        app = LaserStageApp()
        app.withdraw()
        try:
            tab_names = [app.notebook.tab(tab_id, "text") for tab_id in app.notebook.tabs()]
            self.assertIn("Excel Beam Plotter", tab_names)
            app.plotter_dataset = dataset
            app.plotter_canvas.configure(width=800, height=600)
            app.update_idletasks()
            app._redraw_excel_beam_preview()
            self.assertGreater(len(app.plotter_canvas.find_all()), 25)
            for view_name in (
                "Measured point map",
                "X centerline profile",
                "Z centerline profile",
                "Radial average profile",
            ):
                app.plotter_view_var.set(view_name)
                app._redraw_excel_beam_preview()
                self.assertGreater(len(app.plotter_canvas.find_all()), 10)
        finally:
            app.destroy()


if __name__ == "__main__":
    unittest.main()
